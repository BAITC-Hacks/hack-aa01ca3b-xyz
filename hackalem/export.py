"""DOCX and PDF generation for reviewed meeting minutes."""

from __future__ import annotations

import io
import os
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def build_docx(title: str, meeting_date: str, summary: str, actions: list[dict[str, Any]], transcript: list[dict[str, Any]]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.add_heading(title or "Протокол совещания", 0)
    document.add_paragraph(f"Дата совещания: {meeting_date or 'не указана'}")
    document.add_heading("Краткое саммари", level=1)
    document.add_paragraph(summary or "Саммари не сформировано.")
    document.add_heading(f"Поручения ({len(actions)})", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Light Shading Accent 1"
    for cell, heading in zip(table.rows[0].cells, ["Поручение", "Ответственный", "Спикер", "Срок", "Статус"]):
        cell.text = heading
    for action in actions:
        cells = table.add_row().cells
        for cell, value in zip(cells, [action.get("task"), action.get("assignee"), action.get("speaker"), action.get("deadline"), action.get("status")]):
            cell.text = str(value or "")
    document.add_heading("Транскрипт", level=1)
    for segment in transcript:
        start = _stamp(segment.get("start", 0))
        end = _stamp(segment.get("end", 0))
        speaker = segment.get("speaker", "Спикер не определён")
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"[{start}–{end}] {speaker}: ")
        run.bold = True
        paragraph.add_run(segment.get("text", ""))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(title: str, meeting_date: str, summary: str, actions: list[dict[str, Any]], transcript: list[dict[str, Any]]) -> bytes:
    buffer = io.BytesIO()
    font_name, bold_name = _register_unicode_font()
    base_styles = getSampleStyleSheet()
    body = ParagraphStyle("MeetingBody", parent=base_styles["BodyText"], fontName=font_name, fontSize=9, leading=13, alignment=TA_LEFT, spaceAfter=4)
    heading = ParagraphStyle("MeetingHeading", parent=base_styles["Heading2"], fontName=bold_name, fontSize=13, leading=16, spaceBefore=8, spaceAfter=6)
    title_style = ParagraphStyle("MeetingTitle", parent=base_styles["Title"], fontName=bold_name, fontSize=20, leading=24)
    story: list[Any] = [Paragraph(escape(title or "Протокол совещания"), title_style), Spacer(1, 4 * mm)]
    story.append(Paragraph(f"Дата совещания: {escape(meeting_date or 'не указана')}", body))
    story.extend([Paragraph("Краткое саммари", heading), Paragraph(_html(summary or "Саммари не сформировано."), body)])
    story.append(Paragraph(f"Поручения ({len(actions)})", heading))
    rows: list[list[Any]] = [[Paragraph(f"<b>{escape(label)}</b>", body) for label in ("Поручение", "Ответственный", "Спикер", "Срок", "Статус")]]
    for action in actions:
        values = [action.get("task"), action.get("assignee"), action.get("speaker"), action.get("deadline"), action.get("status")]
        rows.append([Paragraph(_html(str(value or "")), body) for value in values])
    table = Table(rows, colWidths=[68 * mm, 29 * mm, 27 * mm, 30 * mm, 24 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c5cfdd")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    story.append(Paragraph("Транскрипт", heading))
    for segment in transcript:
        item = f"<b>[{_stamp(segment.get('start', 0))}–{_stamp(segment.get('end', 0))}] {escape(str(segment.get('speaker', 'Спикер не определён')))}:</b> {_html(segment.get('text', ''))}"
        story.append(Paragraph(item, body))
    SimpleDocTemplate(buffer, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=title or "Протокол совещания").build(story)
    return buffer.getvalue()


def _register_unicode_font() -> tuple[str, str]:
    candidates = [
        (os.getenv("MEETING_FONT"), os.getenv("MEETING_FONT_BOLD")),
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
    ]
    for regular, bold in candidates:
        if regular and Path(regular).is_file():
            pdfmetrics.registerFont(TTFont("MeetingSans", regular))
            if bold and Path(bold).is_file():
                pdfmetrics.registerFont(TTFont("MeetingSansBold", bold))
                return "MeetingSans", "MeetingSansBold"
            return "MeetingSans", "MeetingSans"
    return "Helvetica", "Helvetica-Bold"


def _stamp(seconds: float) -> str:
    total = max(0, int(float(seconds or 0)))
    return f"{total // 60:02d}:{total % 60:02d}"


def _html(value: Any) -> str:
    return escape(str(value or "")).replace("\n", "<br/>")
