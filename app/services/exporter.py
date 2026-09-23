"""DOCX and PDF generation for reviewed meeting minutes."""

from __future__ import annotations

import io
import os
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import FONTS_DIR

def build_docx(
    title: str,
    meeting_date: str,
    summary: str,
    actions: list[dict[str, Any]],
    transcript: list[dict[str, Any]],
    participants: list[dict[str, Any]] | None = None,
    summary_items: list[dict[str, Any]] | None = None,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.add_heading(title or "Протокол совещания", 0)
    document.add_paragraph(f"Дата совещания: {meeting_date or 'не указана'}")
    document.add_paragraph("Участники уведомлены о ведении записи и её расшифровке с помощью ИИ. Обработка выполнена локально.")
    if participants:
        document.add_heading("Участники", level=1)
        for person in participants:
            document.add_paragraph(f"{person.get('speaker', '')}: {person.get('name', '')}", style="List Bullet")
    document.add_heading("Краткое саммари", level=1)
    document.add_paragraph(summary or "Саммари не сформировано.")
    if summary_items:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        for cell, heading in zip(table.rows[0].cells, ["Направление / доклад", "Показатель", "Проблема"]):
            cell.text = heading
        for item in summary_items:
            for cell, key in zip(table.add_row().cells, ("topic", "indicator", "problem")):
                cell.text = str(item.get(key) or "")
    document.add_heading(f"Поручения ({len(actions)})", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    for cell, heading in zip(table.rows[0].cells, ACTION_HEADINGS):
        cell.text = heading
    for action in actions:
        cells = table.add_row().cells
        for cell, value in zip(cells, _action_row(action)):
            cell.text = value
    document.add_heading("Транскрипт", level=1)
    for segment in transcript:
        start = _stamp(segment.get("start", 0))
        end = _stamp(segment.get("end", 0))
        speaker = segment.get("speaker", "Спикер не определён")
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"[{start}–{end}] {speaker}: ")
        run.bold = True
        paragraph.add_run(segment.get("text", ""))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(
    title: str,
    meeting_date: str,
    summary: str,
    actions: list[dict[str, Any]],
    transcript: list[dict[str, Any]],
    participants: list[dict[str, Any]] | None = None,
    summary_items: list[dict[str, Any]] | None = None,
) -> bytes:
    buffer = io.BytesIO()
    font_name, bold_name = _register_unicode_font()
    base_styles = getSampleStyleSheet()
    body = ParagraphStyle("MeetingBody", parent=base_styles["BodyText"], fontName=font_name, fontSize=9, leading=13, alignment=TA_LEFT, spaceAfter=4)
    heading = ParagraphStyle("MeetingHeading", parent=base_styles["Heading2"], fontName=bold_name, fontSize=13, leading=16, spaceBefore=8, spaceAfter=6)
    title_style = ParagraphStyle("MeetingTitle", parent=base_styles["Title"], fontName=bold_name, fontSize=20, leading=24)
    story: list[Any] = [Paragraph(escape(title or "Протокол совещания"), title_style), Spacer(1, 4 * mm)]
    story.append(Paragraph(f"Дата совещания: {escape(meeting_date or 'не указана')}", body))
    story.append(Paragraph("Участники уведомлены о ведении записи и её расшифровке с помощью ИИ. Обработка выполнена локально.", body))
    if participants:
        story.append(Paragraph("Участники", heading))
        for person in participants:
            story.append(Paragraph(_html(f"{person.get('speaker', '')}: {person.get('name', '')}"), body))
    story.extend([Paragraph("Краткое саммари", heading), Paragraph(_html(summary or "Саммари не сформировано."), body)])
    if summary_items:
        summary_rows: list[list[Any]] = [[Paragraph(f"<b>{escape(label)}</b>", body) for label in ("Направление / доклад", "Показатель", "Проблема")]]
        for item in summary_items:
            summary_rows.append([Paragraph(_html(str(item.get(key) or "")), body) for key in ("topic", "indicator", "problem")])
        summary_table = Table(summary_rows, colWidths=[60 * mm, 50 * mm, 68 * mm], repeatRows=1, hAlign="LEFT")
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c5cfdd")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(summary_table)
    story.append(Paragraph(f"Поручения ({len(actions)})", heading))
    rows: list[list[Any]] = [[Paragraph(f"<b>{escape(label)}</b>", body) for label in ACTION_HEADINGS]]
    for action in actions:
        rows.append([Paragraph(_html(value), body) for value in _action_row(action)])
    table = Table(rows, colWidths=[56 * mm, 32 * mm, 26 * mm, 30 * mm, 20 * mm, 14 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c5cfdd")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    story.append(Paragraph("Транскрипт", heading))
    for segment in transcript:
        item = f"<b>[{_stamp(segment.get('start', 0))}–{_stamp(segment.get('end', 0))}] {escape(str(segment.get('speaker', 'Спикер не определён')))}:</b> {_html(segment.get('text', ''))}"
        story.append(Paragraph(item, body))
    SimpleDocTemplate(buffer, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=title or "Протокол совещания").build(story)
    return buffer.getvalue()


ACTION_HEADINGS = ["Поручение", "Ответственный", "Голос", "Срок", "Статус", "Время"]


def _action_row(action: dict[str, Any]) -> list[str]:
    deadline = str(action.get("deadline") or "Не указан")
    if action.get("deadline_iso"):
        deadline = f"{deadline} ({action['deadline_iso']})"
    return [
        str(action.get("task") or ""),
        str(action.get("assignee") or ""),
        str(action.get("speaker") or ""),
        deadline,
        str(action.get("status") or ""),
        str(action.get("time") or ""),
    ]


def _register_unicode_font() -> tuple[str, str]:
    candidates = [
        (os.getenv("MEETING_FONT"), os.getenv("MEETING_FONT_BOLD")),
        (str(FONTS_DIR / "DejaVuSans.ttf"), str(FONTS_DIR / "DejaVuSans-Bold.ttf")),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
    ]
    for regular, bold in candidates:
        if regular and Path(regular).is_file():
            pdfmetrics.registerFont(TTFont("MeetingSans", regular))
            if bold and Path(bold).is_file():
                pdfmetrics.registerFont(TTFont("MeetingSansBold", bold))
                return "MeetingSans", "MeetingSansBold"
            return "MeetingSans", "MeetingSans"
    return "Helvetica", "Helvetica-Bold"


def _stamp(seconds: float) -> str:
    total = max(0, int(float(seconds or 0)))
    return f"{total // 60:02d}:{total % 60:02d}"


def _html(value: Any) -> str:
    return escape(str(value or "")).replace("\n", "<br/>")


def build_transcript_pdf(title: str, meeting_date: str, transcript: list[dict[str, Any]]) -> bytes:
    """Transcript-only PDF (e.g. right after live transcription)."""
    buffer = io.BytesIO()
    font_name, bold_name = _register_unicode_font()
    base_styles = getSampleStyleSheet()
    body = ParagraphStyle("TranscriptBody", parent=base_styles["BodyText"], fontName=font_name, fontSize=10, leading=14, alignment=TA_LEFT, spaceAfter=5)
    title_style = ParagraphStyle("TranscriptTitle", parent=base_styles["Title"], fontName=bold_name, fontSize=18, leading=22)
    story: list[Any] = [Paragraph(escape(title or "Стенограмма совещания"), title_style), Spacer(1, 3 * mm)]
    story.append(Paragraph(f"Дата: {escape(meeting_date or 'не указана')} · Сформировано Briefly AI локально", body))
    story.append(Spacer(1, 3 * mm))
    for segment in transcript:
        speaker = escape(str(segment.get("speaker", "")))
        prefix = f"<b>[{_stamp(segment.get('start', 0))}] {speaker}:</b> " if speaker else f"<b>[{_stamp(segment.get('start', 0))}]</b> "
        story.append(Paragraph(prefix + _html(segment.get("text", "")), body))
    SimpleDocTemplate(buffer, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=title or "Стенограмма").build(story)
    return buffer.getvalue()
